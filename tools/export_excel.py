#!/usr/bin/env python3
"""把 starred_vs_local.json + 本地未 star 清单导出成多 sheet Excel 文件"""
import json, os, re
from datetime import datetime
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
SRC = os.path.join(DATA_DIR, "starred_vs_local.json")
LOCAL_REPOS_TXT = os.path.join(DATA_DIR, "local_repos.txt")
OUT = os.path.join(DATA_DIR, "GitHub_Star_vs_Local_对比报告.xlsx")

CATEGORY_CN = {
    "AI-Agent-Framework": "AI Agent 框架/Harness",
    "AI-Coding-Harness": "AI 编码 Agent / Harness",
    "AI-Agent-Skill": "Agent Skills / 技能库",
    "AI-RAG-Knowledge": "RAG / 知识库 / 向量检索",
    "AI-LLM-Training": "LLM 训练 / 微调 / 推理",
    "AI-LLM-App": "LLM 应用合集 / 提示词",
    "AI-Video": "AI 视频 / 短剧 / 数字人",
    "AI-Voice-TTS": "语音合成 / TTS / 音乐 / 数字人",
    "AI-Vision-Image": "图像生成 / 多模态 / AI 记忆",
    "AI-Agent-Browser": "浏览器 Agent / Web 自动化",
    "AI-Agent-Social": "社媒采集 / 舆情 / Agent 触达",
    "AI-Productivity": "AI 生产力 / 助手 / 办公",
    "AI-Security-RedTeam": "AI 安全 / 红队 / 防护",
    "AI-Data-Infra": "数据工程 / 工作流编排",
    "AI-Book-Tutorial": "Awesome 合集 / 教程书籍 / 文档",
    "Ops-Observability": "运维 / 可观测性 / 分析平台",
    "DevOps-Platform": "DevOps / CI/CD 平台",
    "DevTools-Plugins": "开发工具 / 插件 / MCP",
    "Media-Tool-NonAI": "媒体工具 / 下载（弱 AI）",
    "Misc": "其他 / 未分类",
}

# ---------- 样式 ----------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
CAT_FILL   = PatternFill("solid", fgColor="DDEBF7")
CAT_FONT   = Font(bold=True, color="1F4E78")
AI_YES_FILL = PatternFill("solid", fgColor="FFF2CC")
MONEY_FILL  = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

def style_header(ws, row=1, ncols=None):
    if ncols is None:
        ncols = ws.max_column
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

def auto_col_width(ws, max_w=60):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            # 中文宽度按 2 估
            w = sum(2 if ord(ch) > 127 else 1 for ch in v)
            if w > longest:
                longest = w
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), max_w)

def add_row(ws, values, row=None):
    r = row or ws.max_row + 1
    for i, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.border = BORDER
        cell.alignment = LEFT
    return r

# ---------- 载入数据 ----------
with open(SRC) as f:
    data = json.load(f)

cloned   = data["cloned_items"]
uncloned = sorted(data["not_cloned_items"], key=lambda x: -x.get("stars", 0))

# 本地有 clone 但没 star 的项目
starred_fulls = {x["full_name"].lower().rstrip("/") for x in cloned + uncloned}
starred_names = {x["full_name"].split("/")[-1].lower() for x in cloned + uncloned}
no_star_rows = []
with open(LOCAL_REPOS_TXT) as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        parts = line.split("|")
        if len(parts) < 3:
            continue
        local_dir, repo_full, remote = parts[0], parts[1].rstrip("/"), parts[2]
        key = repo_full.lower() if repo_full else ""
        name_only = (repo_full.split("/")[-1].lower() if repo_full else local_dir.lower())
        if not key or key == "not-a-git-repo":
            if local_dir == "docs":
                no_star_rows.append((local_dir, "", "(本地文档目录，非 GitHub 仓库)"))
            continue
        if key not in starred_fulls and name_only not in starred_names:
            no_star_rows.append((local_dir, repo_full, remote))

# ---------- 生成 workbook ----------
wb = Workbook()

# ============= Sheet 1: 总览 =============
ws = wb.active
ws.title = "总览"
ws["A1"] = f"GitHub Star vs 本地 aiengine/ 对比报告   — 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
ws.merge_cells("A1:E1")
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

ws["A3"] = "📊 核心数字"
ws["A3"].font = Font(bold=True, size=12, color="1F4E78")
overview_kv = [
    ("你在 GitHub 上 star 项目总数", len(cloned) + len(uncloned)),
    ("已克隆到 aiengine/ 的数量", len(cloned)),
    ("star 过但未克隆数量", len(uncloned)),
    ("AI 相关项目（star 列表里）", sum(1 for x in cloned + uncloned if x["_is_ai"])),
    ("已克隆中 AI 相关", sum(1 for x in cloned if x["_is_ai"])),
    ("已克隆中非 AI（弱/工具类）", sum(1 for x in cloned if not x["_is_ai"])),
    ("本地有 clone 但未 star 的项目数", len(no_star_rows)),
]
for i, (k, v) in enumerate(overview_kv, start=4):
    ws.cell(row=i, column=1, value=k).border = BORDER
    cell = ws.cell(row=i, column=2, value=v)
    cell.border = BORDER
    cell.alignment = CENTER
    cell.fill = MONEY_FILL
    cell.font = Font(bold=True)

ws["A13"] = "📂 分类分布（AI 类优先）"
ws["A13"].font = Font(bold=True, size=12, color="1F4E78")
headers_cat = ["分类编码", "分类中文", "已克隆数", "未克隆数", "合计", "AI 相关?"]
for i, h in enumerate(headers_cat, 1):
    ws.cell(row=14, column=i, value=h)
style_header(ws, row=14, ncols=len(headers_cat))

cats_all = list(CATEGORY_CN.keys())
# 先统计
cloned_counter = Counter(x["_category"] for x in cloned)
uncloned_counter = Counter(x["_category"] for x in uncloned)

row = 15
for cat in cats_all:
    cn = CATEGORY_CN.get(cat, cat)
    c1 = cloned_counter.get(cat, 0)
    c2 = uncloned_counter.get(cat, 0)
    ai = "🤖AI" if cat.startswith("AI-") else "🔧工具"
    add_row(ws, [cat, cn, c1, c2, c1 + c2, ai], row=row)
    if ai == "🤖AI":
        for c in range(1, 7):
            ws.cell(row=row, column=c).fill = AI_YES_FILL
    if c1 > 0 and c2 > 0:
        pass
    row += 1

# 合计行
add_row(ws, ["合计", "", len(cloned), len(uncloned), len(cloned)+len(uncloned), ""], row=row)
for c in range(1, 7):
    cell = ws.cell(row=row, column=c)
    cell.fill = CAT_FILL
    cell.font = Font(bold=True)
auto_col_width(ws, max_w=48)

# ============= Sheet 2: 已克隆清单 =============
ws2 = wb.create_sheet(f"✅ 已克隆（{len(cloned)}）")
headers2 = ["分类中文", "AI 相关", "GitHub Stars", "项目（owner/repo）", "本地目录名", "语言", "一句话解决什么问题"]
for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)
style_header(ws2, row=1, ncols=len(headers2))
ws2.freeze_panes = "A2"

# 按分类分组、组内按 star 降序（分类按 AI 优先）
by_cat = defaultdict(list)
for x in cloned:
    by_cat[x["_category"]].append(x)

r = 2
cat_order = [c for c in cats_all if c in by_cat]
for cat in cat_order:
    # 分类分组头
    merge_from = r
    cn = CATEGORY_CN.get(cat, cat)
    ai = "🤖AI" if cat.startswith("AI-") else "🔧工具"
    for x in sorted(by_cat[cat], key=lambda i: -i.get("stars", 0)):
        add_row(ws2, [
            cn, ai,
            x.get("stars", 0),
            x["full_name"],
            x.get("_local_dir") or "",
            x.get("language") or "-",
            x.get("_summary") or "",
        ], row=r)
        if ai == "🤖AI":
            ws2.cell(row=r, column=2).fill = AI_YES_FILL
        r += 1
    # 分类分组着色
    if r - merge_from > 0:
        for rr in range(merge_from, r):
            ws2.cell(row=rr, column=1).fill = CAT_FILL
            ws2.cell(row=rr, column=1).font = CAT_FONT

ws2.column_dimensions["A"].width = 28
ws2.column_dimensions["B"].width = 10
ws2.column_dimensions["C"].width = 12
ws2.column_dimensions["D"].width = 48
ws2.column_dimensions["E"].width = 30
ws2.column_dimensions["F"].width = 14
ws2.column_dimensions["G"].width = 80
# 数字列格式
for rr in range(2, ws2.max_row + 1):
    c = ws2.cell(row=rr, column=3)
    c.number_format = "#,##0"

# ============= Sheet 3: 未克隆清单 =============
ws3 = wb.create_sheet(f"❌ 未克隆（{len(uncloned)}）")
headers3 = ["分类中文", "AI 相关", "GitHub Stars", "项目（owner/repo）", "语言", "一句话解决什么问题", "Topics"]
for i, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=i, value=h)
style_header(ws3, row=1, ncols=len(headers3))
ws3.freeze_panes = "A2"

by_cat_nc = defaultdict(list)
for x in uncloned:
    by_cat_nc[x["_category"]].append(x)

r = 2
for cat in [c for c in cats_all if c in by_cat_nc]:
    cn = CATEGORY_CN.get(cat, cat)
    ai = "🤖AI" if cat.startswith("AI-") else "🔧工具"
    for x in sorted(by_cat_nc[cat], key=lambda i: -i.get("stars", 0)):
        topics = ", ".join(x.get("topics") or [])
        add_row(ws3, [
            cn, ai,
            x.get("stars", 0),
            x["full_name"],
            x.get("language") or "-",
            x.get("_summary") or "",
            topics,
        ], row=r)
        if ai == "🤖AI":
            ws3.cell(row=r, column=2).fill = AI_YES_FILL
        r += 1
    merge_from = r - len(by_cat_nc[cat])
    for rr in range(merge_from, r):
        ws3.cell(row=rr, column=1).fill = CAT_FILL
        ws3.cell(row=rr, column=1).font = CAT_FONT

ws3.column_dimensions["A"].width = 28
ws3.column_dimensions["B"].width = 10
ws3.column_dimensions["C"].width = 12
ws3.column_dimensions["D"].width = 48
ws3.column_dimensions["E"].width = 14
ws3.column_dimensions["F"].width = 80
ws3.column_dimensions["G"].width = 50
for rr in range(2, ws3.max_row + 1):
    ws3.cell(row=rr, column=3).number_format = "#,##0"

# ============= Sheet 4: 本地 clone 未 star =============
ws4 = wb.create_sheet(f"🔎 本地未 star（{len(no_star_rows)}）")
headers4 = ["本地目录名", "对应 GitHub 项目（owner/repo）", "Remote URL / 说明"]
for i, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=i, value=h)
style_header(ws4, row=1, ncols=len(headers4))
ws4.freeze_panes = "A2"
for r, (d, repo, remote) in enumerate(no_star_rows, start=2):
    add_row(ws4, [d, repo or "-", remote], row=r)
    if not repo:
        for c in range(1, 4):
            ws4.cell(row=r, column=c).fill = MONEY_FILL
ws4.column_dimensions["A"].width = 32
ws4.column_dimensions["B"].width = 48
ws4.column_dimensions["C"].width = 70

# ---------- 保存 ----------
wb.save(OUT)
size_kb = os.path.getsize(OUT) / 1024
print(f"✅ Excel 已保存：{OUT}")
print(f"   文件大小：{size_kb:,.1f} KB")
print(f"   Sheet：{[s.title for s in wb.worksheets]}")
print(f"   行数：已克隆={ws2.max_row-1}, 未克隆={ws3.max_row-1}, 本地未star={ws4.max_row-1}")
